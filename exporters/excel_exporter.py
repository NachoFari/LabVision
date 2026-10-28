import os
import time
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

class ExcelExporter:
    """Generador de archivos Excel estilizados con formato corporativo y de laboratorio."""
    @staticmethod
    def export(data_rows, output_path, mode="viscometer", metadata=None):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Datos Laboratorio"

        HEADER_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        TITLE_FILL = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        WARN_FILL = PatternFill(start_color="FEF2F2", end_color="FEF2F2", fill_type="solid")
        
        FONT_TITLE = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
        FONT_SUB = Font(name="Segoe UI", size=9, italic=True, color="94A3B8")
        FONT_HEADER = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        FONT_REGULAR = Font(name="Segoe UI", size=10, color="0F172A")
        FONT_WARN = Font(name="Segoe UI", size=10, color="DC2626", bold=True)
        
        BORDER_THIN = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        title_text = "REGISTRO VISCOSÍMETRO NDJ-8S" if mode == "viscometer" else "REGISTRO BALANZA ANALÍTICA DE PRECISIÓN"
        last_col = "H" if mode == "viscometer" else "F"

        # Fila 1: Título Principal
        ws.merge_cells(f"A1:{last_col}1")
        cell_title = ws["A1"]
        cell_title.value = f"🔬 {title_text} • LABVISION PRO"
        cell_title.font = FONT_TITLE
        cell_title.fill = TITLE_FILL
        cell_title.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 35

        # Fila 2: Subtítulo
        sub_text = f"Generado automáticamente • Creado por NachoFari • Fecha: {time.strftime('%Y-%m-%d %H:%M:%S')}"
        ws.merge_cells(f"A2:{last_col}2")
        cell_sub = ws["A2"]
        cell_sub.value = sub_text
        cell_sub.font = FONT_SUB
        cell_sub.fill = TITLE_FILL
        cell_sub.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 20

        # Encabezados con columna de Confianza (%)
        if mode == "viscometer":
            headers = [
                ("Tiempo (s)", 14),
                ("Timestamp", 14),
                ("Rotor (#)", 12),
                ("Velocidad (RPM)", 18),
                ("Viscosidad (mPa·s)", 22),
                ("Torque (%)", 14),
                ("Confianza (%)", 15),
                ("Calidad Lectura", 22)
            ]
        else:
            unit = metadata.get("unit", "g") if metadata else "g"
            headers = [
                ("Tiempo (s)", 14),
                ("Timestamp", 14),
                (f"Peso ({unit})", 18),
                ("Unidad", 12),
                ("Confianza (%)", 15),
                ("Calidad Lectura", 22)
            ]

        header_row = 4
        for col_idx, (h_name, width) in enumerate(headers, start=1):
            cell = ws.cell(row=header_row, column=col_idx, value=h_name)
            cell.font = FONT_HEADER
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = BORDER_THIN
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = width

        ws.row_dimensions[header_row].height = 28

        # Filas de datos
        start_data_row = 5
        for i, row in enumerate(data_rows):
            current_row = start_data_row + i
            ws.row_dimensions[current_row].height = 20
            is_bad = row.get("is_bad", False)
            row_fill = WARN_FILL if is_bad else None

            # Calcular porcentaje de confianza legible
            conf_val = row.get("confidence")
            if conf_val is not None:
                conf_pct_str = f"{round(float(conf_val) * 100, 1)}%" if float(conf_val) <= 1.0 else f"{round(float(conf_val), 1)}%"
            else:
                conf_pct_str = "-"

            if mode == "viscometer":
                values = [
                    round(row.get("time_sec", i), 2),
                    row.get("timestamp_str", f"{i}s"),
                    row.get("rotor", "-"),
                    row.get("speed_rpm", "-"),
                    row.get("viscosity_mpa", "-"),
                    row.get("torque_pct", "-"),
                    conf_pct_str,
                    "⚠️ DUDOSA" if is_bad else "✅ ÓPTIMA"
                ]
            else:
                values = [
                    round(row.get("time_sec", i), 2),
                    row.get("timestamp_str", f"{i}s"),
                    row.get("weight", "-"),
                    row.get("unit", "g"),
                    conf_pct_str,
                    "⚠️ DUDOSA" if is_bad else "✅ ÓPTIMA"
                ]

            for col_idx, val in enumerate(values, start=1):
                cell = ws.cell(row=current_row, column=col_idx, value=val)
                cell.font = FONT_WARN if (is_bad and col_idx == len(values)) else FONT_REGULAR
                if row_fill:
                    cell.fill = row_fill
                cell.border = BORDER_THIN
                if isinstance(val, (int, float)):
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="center", vertical="center")

        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        wb.save(output_path)
        print(f"[LabVision] Archivo Excel guardado en: {output_path}")
        return output_path
